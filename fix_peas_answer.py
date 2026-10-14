from docx import Document

docx_path = '/Users/posty/Documents/Arthur-and-Germaine/Arthur_and_Germaine_editable.docx'
doc = Document(docx_path)

found = False
for i in range(len(doc.paragraphs) - 2):
    # Look for "Peas?" then "— Arthur —" then "No."
    if "Peas?" in doc.paragraphs[i].text and "— Arthur —" in doc.paragraphs[i+1].text and "No." in doc.paragraphs[i+2].text:
        doc.paragraphs[i+2].text = "Fine."
        print(f"Fixed Peas answer at paragraph {i+2}")
        found = True
        break

if found:
    doc.save(docx_path)
    print("Saved changes.")
else:
    print("Could not find the Peas sequence.")
