from docx import Document

docx_path = '/Users/posty/Documents/Arthur-and-Germaine/Arthur_and_Germaine_editable.docx'
doc = Document(docx_path)

# We look for the specific sequence: "— Arthur —" followed by "I'm exercising."
# We will merge them into one paragraph.

found = False
for i in range(len(doc.paragraphs) - 1):
    para = doc.paragraphs[i]
    next_para = doc.paragraphs[i+1]
    
    if "— Arthur —" in para.text and "I'm exercising." in next_para.text:
        # Merge text: "— Arthur —\nI'm exercising."
        # To add a newline in docx, we add a run with a break.
        
        # Create the new text for the first paragraph
        para.text = para.text + "\n" + next_para.text
        
        # Delete the second paragraph
        p = next_para._element
        p.getparent().remove(p)
        
        print(f"Merged paragraphs at index {i} and {i+1}")
        found = True
        # Since we modified the list, we should break or adjust index.
        # For this book, let's just do the first one found or all of them.
        # The user said "P.24", so it's likely one specific spot.

if found:
    doc.save(docx_path)
    print("Saved changes.")
else:
    print("Could not find the dialogue split.")
