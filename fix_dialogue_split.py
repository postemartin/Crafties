from docx import Document

doc = Document('/Users/posty/Documents/Arthur-and-Germaine/Arthur_and_Germaine_editable.docx')
changed = False

for i in range(len(doc.paragraphs) - 1):
    para = doc.paragraphs[i]
    next_para = doc.paragraphs[i+1]
    
    if "— Arthur —" in para.text and "I'm exercising." in next_para.text:
        para.paragraph_format.keep_with_next = True
        print(f"Set keep_with_next for para {i} ('{para.text}')")
        changed = True

if changed:
    doc.save('/Users/posty/Documents/Arthur-and-Germaine/Arthur_and_Germaine_editable.docx')
    print("Saved changes.")
else:
    print("No matches found.")
